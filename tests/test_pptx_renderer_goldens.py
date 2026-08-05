import asyncio
import sys
from io import BytesIO
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.models.lesson_schema import (
    BloomLevel,
    ContentBlock,
    ContentMode,
    LearningObjective,
    LearningStructure,
    LessonDeck,
    LessonMetadata,
    PedagogicalRole,
    Slide,
    SlideType,
)
from app.services.pptx_renderer import PPTXRenderer


def _golden_deck(theme: str) -> LessonDeck:
    return LessonDeck(
        meta=LessonMetadata(
            topic=f"{theme.title()} Motion",
            subject="Physics",
            grade="9",
            theme=theme,
            pedagogical_flow="default_classroom",
        ),
        structure=LearningStructure(
            learning_objectives=[
                LearningObjective(objective="Explain motion using force.", bloom_level=BloomLevel.UNDERSTAND),
            ],
            vocabulary=[],
            prerequisites=[],
            bloom_progression=[],
        ),
        slides=[
            Slide(
                title="Core Idea",
                content="Force changes motion.",
                order=1,
                slideType=SlideType.CONCEPT,
                pedagogicalRole=PedagogicalRole.EXPLAIN_CORE,
                contentMode=ContentMode.TEXT_ONLY,
                instructionalGoal="Explain how force changes motion.",
                contentBlocks=[
                    ContentBlock(type="headline", text="Force changes motion"),
                    ContentBlock(type="bullets", items=["Balanced forces keep motion steady", "Unbalanced forces change speed or direction"]),
                ],
            ),
            Slide(
                title="Try It",
                content="Question: What happens when one force is stronger?",
                order=2,
                slideType=SlideType.ACTIVITY,
                pedagogicalRole=PedagogicalRole.GUIDED_PRACTICE,
                contentMode=ContentMode.QUESTION,
                contentBlocks=[
                    ContentBlock(type="question", text="What happens when one force is stronger?"),
                    ContentBlock(type="hint", text="Compare the sizes of the forces."),
                    ContentBlock(type="answer", text="The object accelerates in the stronger force direction."),
                ],
            ),
            Slide(
                title="Wrap Up",
                content="Remember the key ideas.",
                order=3,
                slideType=SlideType.SUMMARY,
                pedagogicalRole=PedagogicalRole.SUMMARY,
                contentMode=ContentMode.TEXT_ONLY,
                contentBlocks=[
                    ContentBlock(type="takeaway", text="Balanced forces do not change motion."),
                    ContentBlock(type="takeaway", text="Unbalanced forces cause acceleration."),
                    ContentBlock(type="takeaway", text="Direction matters as much as size."),
                ],
            ),
        ],
    )


def test_golden_theme_fixtures_render_to_valid_presentations():
    for theme in ("default", "blueprint", "deep"):
        renderer = PPTXRenderer(theme=theme)
        output = asyncio.run(renderer.render_lesson_deck(_golden_deck(theme)))

        assert isinstance(output, BytesIO)
        payload = output.getvalue()
        assert len(payload) > 5000

        prs = Presentation(BytesIO(payload))
        assert len(prs.slides) == 5


def _phase3_flow_deck(theme: str, topic: str, subject: str, flow: str) -> LessonDeck:
    if flow == "problem_solving_math":
        return LessonDeck(
            meta=LessonMetadata(
                topic=topic,
                subject=subject,
                grade="10",
                theme=theme,
                pedagogical_flow=flow,
            ),
            structure=LearningStructure(
                learning_objectives=[
                    LearningObjective(objective="Solve quadratic equations using structured steps.", bloom_level=BloomLevel.APPLY),
                ],
                vocabulary=[],
                prerequisites=["Linear equations", "Factoring basics"],
                bloom_progression=[BloomLevel.UNDERSTAND, BloomLevel.APPLY, BloomLevel.APPLY, BloomLevel.ANALYZE],
            ),
            slides=[
                Slide(
                    title="Start With the Pattern",
                    content="A worked example reveals the solving path.",
                    order=1,
                    slideType=SlideType.INTRODUCTION,
                    pedagogicalRole=PedagogicalRole.HOOK,
                    contentMode=ContentMode.TEXT_ONLY,
                    contentBlocks=[
                        ContentBlock(type="headline", text="Notice the equation pattern"),
                        ContentBlock(type="bullets", items=["Look for factors first", "Track the sign of each term"]),
                    ],
                ),
                Slide(
                    title="Worked Example",
                    content="Factor and solve x^2 + 5x + 6 = 0.",
                    order=2,
                    slideType=SlideType.CONCEPT,
                    pedagogicalRole=PedagogicalRole.WORKED_EXAMPLE,
                    contentMode=ContentMode.EQUATION,
                    instructionalGoal="Show a complete worked solution before abstract explanation.",
                    contentBlocks=[
                        ContentBlock(type="equation", value="x^2 + 5x + 6 = 0"),
                        ContentBlock(type="steps", items=["Find two numbers that sum to 5 and multiply to 6", "Rewrite as (x + 2)(x + 3) = 0", "Set each factor to zero"]),
                    ],
                ),
                Slide(
                    title="Guided Solve",
                    content="Students complete the next equation with prompts.",
                    order=3,
                    slideType=SlideType.ACTIVITY,
                    pedagogicalRole=PedagogicalRole.GUIDED_PRACTICE,
                    contentMode=ContentMode.QUESTION,
                    contentBlocks=[
                        ContentBlock(type="question", text="Solve x^2 + 7x + 12 = 0"),
                        ContentBlock(type="hint", text="Use factor pairs for 12."),
                        ContentBlock(type="answer", text="x = -3 or x = -4"),
                    ],
                ),
                Slide(
                    title="Why the Method Works",
                    content="Connect factoring to the zero-product rule.",
                    order=4,
                    slideType=SlideType.SUMMARY,
                    pedagogicalRole=PedagogicalRole.SUMMARY,
                    contentMode=ContentMode.TEXT_ONLY,
                    contentBlocks=[
                        ContentBlock(type="takeaway", text="Factoring turns one equation into two simpler equations."),
                        ContentBlock(type="takeaway", text="The zero-product rule explains why each factor can be set to zero."),
                    ],
                ),
            ],
        )

    return LessonDeck(
        meta=LessonMetadata(
            topic=topic,
            subject=subject,
            grade="10",
            theme=theme,
            pedagogical_flow=flow,
        ),
        structure=LearningStructure(
            learning_objectives=[
                LearningObjective(objective="Recall and connect the stages of cell division.", bloom_level=BloomLevel.UNDERSTAND),
            ],
            vocabulary=[],
            prerequisites=["Cell structure"],
            bloom_progression=[BloomLevel.REMEMBER, BloomLevel.UNDERSTAND, BloomLevel.APPLY, BloomLevel.APPLY],
        ),
        slides=[
            Slide(
                title="Fast Recall",
                content="Review the order of mitosis stages.",
                order=1,
                slideType=SlideType.INTRODUCTION,
                pedagogicalRole=PedagogicalRole.HOOK,
                contentMode=ContentMode.TEXT_ONLY,
                contentBlocks=[
                    ContentBlock(type="headline", text="What do you already remember?"),
                    ContentBlock(type="bullets", items=["Prophase", "Metaphase", "Anaphase", "Telophase"]),
                ],
            ),
            Slide(
                title="Core Refresh",
                content="Summarize what changes in each stage.",
                order=2,
                slideType=SlideType.CONCEPT,
                pedagogicalRole=PedagogicalRole.EXPLAIN_CORE,
                contentMode=ContentMode.DIAGRAM,
                instructionalGoal="Refresh the stage-by-stage logic before revision practice.",
                contentBlocks=[
                    ContentBlock(type="headline", text="Stage-by-stage review"),
                    ContentBlock(type="bullets", items=["Chromosomes condense", "Chromosomes line up", "Chromatids separate"]),
                ],
            ),
            Slide(
                title="Revision Check 1",
                content="Students identify the correct stage from clues.",
                order=3,
                slideType=SlideType.ACTIVITY,
                pedagogicalRole=PedagogicalRole.GUIDED_PRACTICE,
                contentMode=ContentMode.QUESTION,
                contentBlocks=[
                    ContentBlock(type="question", text="Which stage shows chromosomes lined up at the center?"),
                    ContentBlock(type="answer", text="Metaphase"),
                ],
            ),
            Slide(
                title="Revision Check 2",
                content="Students compare two similar stage descriptions.",
                order=4,
                slideType=SlideType.ACTIVITY,
                pedagogicalRole=PedagogicalRole.GUIDED_PRACTICE,
                contentMode=ContentMode.COMPARISON,
                contentBlocks=[
                    ContentBlock(type="question", text="How is anaphase different from telophase?"),
                    ContentBlock(type="answer", text="Anaphase separates chromatids; telophase reforms nuclei."),
                ],
            ),
            Slide(
                title="Exam Wrap",
                content="Condense the revision into three takeaways.",
                order=5,
                slideType=SlideType.SUMMARY,
                pedagogicalRole=PedagogicalRole.SUMMARY,
                contentMode=ContentMode.TEXT_ONLY,
                contentBlocks=[
                    ContentBlock(type="takeaway", text="Remember the order of stages."),
                    ContentBlock(type="takeaway", text="Link each stage to one visible chromosome event."),
                    ContentBlock(type="takeaway", text="Use comparisons to avoid common exam confusion."),
                ],
            ),
        ],
    )


def test_problem_solving_and_revision_fixtures_render_successfully():
    scenarios = [
        ("mathematics", "Quadratic Equations", "Mathematics", "problem_solving_math"),
        ("science_nature", "Cell Division", "Biology", "revision_exam_prep"),
        ("rust", "Cell Division", "Biology", "revision_exam_prep"),
    ]

    for theme, topic, subject, flow in scenarios:
        renderer = PPTXRenderer(theme=theme)
        output = asyncio.run(renderer.render_lesson_deck(_phase3_flow_deck(theme, topic, subject, flow)))

        assert isinstance(output, BytesIO)
        payload = output.getvalue()
        assert len(payload) > 5000

        prs = Presentation(BytesIO(payload))
        assert len(prs.slides) >= 6
